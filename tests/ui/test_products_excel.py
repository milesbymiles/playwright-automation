import pytest
from openpyxl import load_workbook

from utils.excel_util import read_sheet, write_sheet

pytestmark = [pytest.mark.ui, pytest.mark.regression]


def products_on_page(inventory_page) -> list[dict]:
    names = inventory_page.get_product_names()
    prices = [float(p.replace("$", "")) for p in inventory_page.get_product_prices()]
    return [{"product": n, "price": p} for n, p in zip(names, prices)]


def test_product_prices_match_the_excel_sheet(logged_in, inventory_page):
    expected = {row["product"]: row["price"] for row in read_sheet("products.xlsx", "prices")}
    actual = {row["product"]: row["price"] for row in products_on_page(inventory_page)}

    assert actual == expected


def test_export_the_product_list_to_an_excel_file(logged_in, inventory_page):
    path = write_sheet("reports/exports/products_export.xlsx", "products", products_on_page(inventory_page))

    data_rows = load_workbook(path)["products"].max_row - 1   # minus the header row
    assert data_rows == 6
