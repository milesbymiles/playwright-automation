from pytest_bdd import when, then, parsers

from openpyxl import load_workbook

from framework.utils.excel_util import read_sheet, write_sheet

EXPORT_DIR = "reports/exports"


def _products_on_page(inventory_page) -> list[dict]:
    names = inventory_page.get_product_names()
    prices = [float(p.replace("$", "")) for p in inventory_page.get_product_prices()]
    return [{"product": n, "price": p} for n, p in zip(names, prices)]


@then(parsers.parse('the product prices match the "{sheet}" sheet in "{file_name}"'))
def prices_match_excel(inventory_page, sheet, file_name):
    expected = {row["product"]: row["price"] for row in read_sheet(file_name, sheet)}
    actual = {row["product"]: row["price"] for row in _products_on_page(inventory_page)}
    assert actual == expected


@when(parsers.parse('the user exports the product list to "{file_name}"'), target_fixture="export_path")
def export_products(inventory_page, file_name):
    return write_sheet(f"{EXPORT_DIR}/{file_name}", "products", _products_on_page(inventory_page))


@then(parsers.parse("the exported file has {count:d} products"))
def exported_file_has_products(export_path, count):
    rows = load_workbook(export_path)["products"].max_row - 1   # minus the header row
    assert rows == count
