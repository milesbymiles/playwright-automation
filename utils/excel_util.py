"""Simple Excel helper. First row of a sheet = column names."""
from pathlib import Path

from openpyxl import Workbook, load_workbook

from utils.data_loader import DATA_DIR


def read_sheet(file_name: str, sheet_name: str) -> list[dict]:
    """Read data/test_data/<file_name>. Returns one dict per row, e.g. {"product": "...", "price": 29.99}"""
    workbook = load_workbook(DATA_DIR / "test_data" / file_name, data_only=True)
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    headers = rows[0]
    return [dict(zip(headers, row)) for row in rows[1:]]


def write_sheet(file_path: str, sheet_name: str, rows: list[dict]) -> Path:
    """Write a list of dicts to an Excel file (headers come from the first dict's keys)."""
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(rows[0].keys()))
    for row in rows:
        sheet.append(list(row.values()))
    workbook.save(path)
    return path
